"""读取本地排单工作簿（只读，不修改文件）。"""

from __future__ import annotations

import os
from collections.abc import Iterable
from pathlib import Path
from typing import Any, Callable

from app.wps.common import FIRST_DATA_ROW, WEEKDAYS, normalize_phone
from app.wps.errors import WpsCloudError
from app.wps.models import CloudOrder


LOCAL_SHEETS = ("东湖中餐", "衣锦中餐", "医学院中餐",
                "东湖晚餐", "衣锦晚餐", "医学院晚餐")

LOCAL_COL = {
    "order": 1, "name": 2, "address": 3, "phone": 4,
    # 周一~周日 七列（1-based）：记录的是**这一批要送的那天**是周几，
    # 云同步用它核对"本地表是不是这次目标日期的批次"（见 planner）。
    "weekdays": (5, 6, 7, 8, 9, 10, 11),
    "type": 12, "kind": 13, "meals": 14,
}


def _weekday_marks(cells: Iterable[Any]) -> tuple[str, ...]:
    """从一行的「周一~周日」七格里读出标记（只为真值格；``0``/空格都不算）。"""
    values = list(cells)
    marks: list[str] = []
    for name, raw in zip(WEEKDAYS, values):
        text = str(raw if raw is not None else "").strip()
        if text and text != "0":
            marks.append(name)
    return tuple(marks)


def read_local_orders(excel_path: str | os.PathLike[str], *,
                      sheets: Iterable[str] = LOCAL_SHEETS,
                      log: Callable[[str], Any] | None = None) -> dict[str, list[CloudOrder]]:
    """读取本地排单工作簿，返回 {子表名: [CloudOrder, ...]}（**一行一条**）。

    同一个人的多行**在这里保持原样**（不合并不去重）：由 ``planner`` 按
    "每行算 1 餐"相加并给出提示（见 ``planner._sum_rows_per_person``）。

    只读取，绝不修改本地文件。
    """
    from openpyxl import load_workbook

    path = Path(excel_path)
    if not path.is_file():
        raise WpsCloudError(f"本地排单表不存在：{path}")

    result: dict[str, list[CloudOrder]] = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in sheets:
            if sheet not in wb.sheetnames:
                if log:
                    log(f"[云同步] 本地表缺少子表「{sheet}」，跳过")
                continue
            ws = wb[sheet]
            orders: list[CloudOrder] = []
            # read_only 模式下 max_row 可能是 None，直接按行迭代最稳。
            for row_idx, cells in enumerate(
                    ws.iter_rows(min_row=FIRST_DATA_ROW, max_col=LOCAL_COL["meals"],
                                 values_only=True), start=FIRST_DATA_ROW):
                name = cells[LOCAL_COL["name"] - 1]
                if name is None or str(name).strip() == "":
                    continue
                meals_raw = cells[LOCAL_COL["meals"] - 1]
                try:
                    meals = int(float(meals_raw)) if meals_raw not in (None, "") else 0
                except (TypeError, ValueError):
                    meals = 0
                orders.append(CloudOrder(
                    sheet=sheet,
                    name=str(name).strip(),
                    address=str(cells[LOCAL_COL["address"] - 1] or "").strip(),
                    phone=normalize_phone(cells[LOCAL_COL["phone"] - 1]),
                    meal_type=str(cells[LOCAL_COL["type"] - 1] or "").strip(),
                    meal_kind=str(cells[LOCAL_COL["kind"] - 1] or "").strip(),
                    meals=meals,
                    row=row_idx,
                    order_no=str(cells[LOCAL_COL["order"] - 1] or "").strip(),
                    rows=(row_idx,),
                    weekday_marks=_weekday_marks(
                        cells[col - 1] for col in LOCAL_COL["weekdays"]),
                ))
            result[sheet] = orders
    finally:
        wb.close()
    return result
