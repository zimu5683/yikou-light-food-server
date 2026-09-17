"""把当天名单留档写回《闪时送.xlsx》。"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any, Callable

from app.ordering.import_models import (
    COL_ADDRESS, COL_NAME, COL_PHONE, DATE_COL, DATE_ROW, FIRST_DATA_ROW,
    MealImport,
)
from app.ordering.roster import _log









def archive_to_excel(excel_path: str | os.PathLike[str], meals: list[MealImport],
                     log: Callable[[str], Any] | None = None) -> dict[str, Any]:
    """把当天名单写进《闪时送.xlsx》留档：清空 A:C 第 3 行起，E1 写日期原文。

    只写要下单的人（不含大西/小）；被跳过的那一餐清空名单并把 E1 清空。
    写入走同目录临时文件 + ``os.replace``，避免中途失败留下半份文件。
    """
    from openpyxl import load_workbook

    path = Path(excel_path)
    if not path.is_file():
        raise FileNotFoundError(f"留档 Excel 不存在：{path}")
    workbook = load_workbook(path)
    try:
        summary: dict[str, Any] = {}
        for meal in meals:
            if meal.meal in workbook.sheetnames:
                sheet = workbook[meal.meal]
            else:
                sheet = workbook.create_sheet(title=meal.meal)
                _log(log, f"留档：工作簿里没有「{meal.meal}」表，已新建", "WARN")
            last_row = max(int(sheet.max_row or 0), FIRST_DATA_ROW)
            for row in range(FIRST_DATA_ROW, last_row + 1):
                for col in (COL_NAME, COL_ADDRESS, COL_PHONE):
                    sheet.cell(row=row, column=col).value = None
            cleared = max(0, last_row - FIRST_DATA_ROW + 1)
            if meal.skipped:
                sheet.cell(row=DATE_ROW, column=DATE_COL).value = None
                summary[meal.meal] = {"written": 0, "cleared": cleared, "date_text": ""}
                continue
            for offset, order in enumerate(meal.orders):
                row = FIRST_DATA_ROW + offset
                sheet.cell(row=row, column=COL_NAME).value = order["name"]
                sheet.cell(row=row, column=COL_ADDRESS).value = order["door"]
                phone_cell = sheet.cell(row=row, column=COL_PHONE)
                phone_cell.value = str(order["phone"])
                # 电话按文本写，避免 Excel 把 11 位号码显示成科学计数。
                phone_cell.number_format = "@"
            sheet.cell(row=DATE_ROW, column=DATE_COL).value = meal.date_text
            summary[meal.meal] = {"written": meal.order_count, "cleared": cleared,
                                  "date_text": meal.date_text}

        handle, temp_name = tempfile.mkstemp(
            prefix=f".{path.stem}.", suffix=".xlsx", dir=str(path.parent))
        os.close(handle)
        temp_path = Path(temp_name)
        try:
            workbook.save(temp_path)
            os.replace(temp_path, path)
        except BaseException:
            temp_path.unlink(missing_ok=True)
            raise
        return summary
    except PermissionError as exc:
        raise PermissionError(
            f"《{path.name}》无法写入（可能正被 Excel/WPS 打开）：{exc}") from exc
    finally:
        workbook.close()
