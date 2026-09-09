"""Tests for the 闪时送 (sss) order placement module."""
from __future__ import annotations

import datetime as dt

import pytest

from app import sss
from app.locators import SSS_LOCATORS, load_sss_locators, sss_user_locators_path


def test_compute_delivery_time_lunch_before_16():
    now = dt.datetime(2026, 8, 14, 9, 30)
    assert sss.compute_delivery_time(False, now) == "2026-08-14 11:00:00"


def test_compute_delivery_time_dinner_before_16():
    now = dt.datetime(2026, 8, 14, 9, 30)
    assert sss.compute_delivery_time(True, now) == "2026-08-14 17:00:00"


def test_compute_delivery_time_rolls_to_next_day_after_16():
    now = dt.datetime(2026, 8, 14, 18, 0)
    assert sss.compute_delivery_time(False, now) == "2026-08-15 11:00:00"
    assert sss.compute_delivery_time(True, now) == "2026-08-15 17:00:00"


def test_compute_delivery_time_rolls_to_next_day_between_16_and_20():
    now = dt.datetime(2026, 8, 14, 17, 0)
    assert sss.compute_delivery_time(False, now) == "2026-08-15 11:00:00"


def test_load_sss_orders_reads_two_sheets_and_stops_on_blank_row(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    lunch = wb.active
    lunch.title = "午餐"
    dinner = wb.create_sheet("晚餐")
    lunch.append(["姓名", "门牌号", "电话", "送达时间"])  # 第 1 行表头
    lunch.append(["表头占位", "", "", ""])                 # 第 2 行占位
    lunch.append(["张三", "A101", "13800000001", "11:00"])  # 第 3 行
    lunch.append(["李四", "B202", "13800000002", "11:00"])  # 第 4 行
    lunch.append([None, None, None, None])                 # 空行×3 才终止
    lunch.append([None, None, None, None])
    lunch.append([None, None, None, None])
    lunch.append(["不应读取", "C303", "13800000003", "11:00"])
    dinner.append(["姓名", "门牌号", "电话", "送达时间"])
    dinner.append(["表头占位", "", "", ""])
    dinner.append(["王五", "D404", "13800000004", "17:00"])

    path = tmp_path / "闪时送.xlsx"
    wb.save(path)
    wb.close()

    result = sss.load_sss_orders(path)
    assert set(result) == {"午餐", "晚餐"}
    assert len(result["午餐"]) == 2
    assert result["午餐"][0]["name"] == "张三"
    assert result["午餐"][0]["door"] == "A101"
    assert result["午餐"][0]["phone"] == "13800000001"
    assert result["午餐"][0]["row"] == 3
    assert result["午餐"][1]["name"] == "李四"
    assert len(result["晚餐"]) == 1
    assert result["晚餐"][0]["name"] == "王五"


def test_load_sss_orders_tolerates_single_blank_row(tmp_path):
    """中间单个空行不再截断：连续 3 个空行才结束。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "午餐"
    ws.append(["姓名", "门牌号", "电话", "送达时间"])
    ws.append(["表头占位", "", "", ""])
    ws.append(["张三", "A101", "13800000001", "11:00"])
    ws.append([None, None, None, None])
    ws.append(["李四", "B202", "13800000002", "11:00"])
    ws.append([None, None, None, None])
    ws.append([None, None, None, None])
    ws.append([None, None, None, None])
    ws.append(["不应读取", "C303", "13800000003", "11:00"])
    path = tmp_path / "闪时送.xlsx"
    wb.save(path)
    wb.close()

    result = sss.load_sss_orders(path)
    assert [o["name"] for o in result["午餐"]] == ["张三", "李四"]


def test_validate_sss_orders_normalises_phone_and_rejects_invalid_rows():
    valid = {"午餐": [{
        "row": 3, "name": " 张三 ", "door": " Ａ 1 ",
        "phone": "１３８ ００００ ０００１",
    }]}
    sss._validate_sss_orders(valid)
    assert valid["午餐"][0] == {
        "row": 3, "name": "张三", "door": "A 1", "phone": "13800000001",
    }
    assert sss._normalise_phone(13800000001.0) == "13800000001"

    invalid = {"午餐": [
        {"row": 3, "name": "", "door": "A1", "phone": "13800000001"},
        {"row": 4, "name": "李四", "door": "", "phone": "13800000001.5"},
    ]}
    with pytest.raises(ValueError, match=r"午餐第 3 行.*午餐第 4 行"):
        sss._validate_sss_orders(invalid)


def test_dry_run_validates_all_excel_rows_before_any_network_path(tmp_path):
    from openpyxl import Workbook
    from types import SimpleNamespace

    wb = Workbook()
    ws = wb.active
    ws.title = "午餐"
    ws.append(["姓名", "门牌号", "电话"])
    ws.append(["表头占位", "", ""])
    ws.append(["张三", "A101", "not-a-phone"])
    path = tmp_path / "闪时送.xlsx"
    wb.save(path)
    wb.close()

    with pytest.raises(ValueError, match=r"午餐第 3 行.*11 位电话"):
        sss.run_sss_job(SimpleNamespace(sss_excel_path=str(path), sss_dry_run=True),
                        object())


def test_substitute_tokens_replaces_placeholders_in_candidates():
    step = {"candidates": [{"text": "{label}"}, {"css": ".x", "has_text": "{label}"}]}
    out = sss._substitute_tokens(step, label="嗯哼")
    assert out["candidates"][0]["text"] == "嗯哼"
    assert out["candidates"][1]["has_text"] == "嗯哼"


def test_sss_locators_are_separate_from_default_table(monkeypatch, tmp_path):
    from app import locators as locators_mod

    monkeypatch.setattr(locators_mod, "user_data_dir", lambda: tmp_path)
    assert sss_user_locators_path() == tmp_path / "sss_locators.json"
    table = load_sss_locators()
    assert table["创建订单"]["candidates"]
    assert table["地址选项"]["candidates"][0]["has_text"] == "{label}"
    # 默认表（管理后台）与闪时送表互不污染。
    assert "门店地址" not in table
    assert "创建订单" not in locators_mod.DEFAULT_LOCATORS


def test_sss_locators_include_required_steps():
    for step in ("创建订单", "预约单选项", "一口轻食选项", "送达时间输入",
                 "顾客姓名", "顾客电话", "门牌号", "最终确定"):
        assert step in SSS_LOCATORS
        assert SSS_LOCATORS[step]["candidates"]


def test_build_order_payload_matches_captured_schema():
    import datetime as dt

    order = {"row": 3, "name": "张三", "door": "A101", "phone": "13800000001"}
    address = {"lnt": 119.727873, "lat": 30.257483,
               "areaCode": "330112", "addressDetail": "浙江农林大学东湖校区"}
    payload = sss.build_order_payload(order, False, 211053, address, "轻食",
                                      now=dt.datetime(2026, 9, 5, 18, 0))
    assert payload == {
        "expectedDeliveryTime": "2026-09-06 11:00:00",  # 18 点后顺延次日
        "goodsDetail": [{"goodsName": "轻食", "goodsNum": 1}],
        "orderType": 2,
        "receiveName": "张三",
        "receivePhone": "13800000001",
        "storeId": 211053,
        "receiveAddress": {
            "lnt": 119.727873, "lat": 30.257483,
            "areaCode": "330112",
            "addressDetail": "浙江农林大学东湖校区",
            "doorNum": "A101",
        },
    }


def test_match_record_prefers_keyword_and_single_record_fallback():
    records = [{"id": 1, "name": "其他"}, {"id": 2, "name": "一口轻食"}]
    assert sss._match_record(records, "一口轻食", "门店")["id"] == 2
    assert sss._match_record([{"id": 9}], "", "门店")["id"] == 9


def test_address_from_record_reads_marklnglat_like_frontend():
    record = {"name": "嗯哼", "markLnglat": {
        "longitude": 119.727873, "latitude": 30.257483,
        "adcode": "330112", "address": "浙江省杭州市临安区浙江农林大学东湖校区"}}
    addr = sss._address_from_record(record)
    assert addr == {"lnt": 119.727873, "lat": 30.257483,
                    "areaCode": "330112",
                    "addressDetail": "浙江省杭州市临安区浙江农林大学东湖校区"}


def test_address_from_record_reads_real_frequent_address_fields():
    """2026-09-05 实测的「嗯哼」记录结构：code=区划码、position=地址文本。"""
    record = {"id": 75812, "longitude": "119.728224", "latitude": "30.256632",
              "code": "330110", "position": "浙江农林大学东湖校区",
              "houseNum": "图书馆", "contactName": "嗯哼"}
    addr = sss._address_from_record(record)
    assert addr == {"lnt": 119.728224, "lat": 30.256632,
                    "areaCode": "330110", "addressDetail": "浙江农林大学东湖校区"}


def test_fixed_address_from_config_reads_config():
    from types import SimpleNamespace

    cfg = SimpleNamespace(
        sss_fixed_lnt="119.728224",
        sss_fixed_lat="30.256632",
        sss_fixed_area_code="330110",
        sss_fixed_address_detail="浙江农林大学东湖校区",
    )
    assert sss._fixed_address_from_config(cfg) == {
        "lnt": 119.728224,
        "lat": 30.256632,
        "areaCode": "330110",
        "addressDetail": "浙江农林大学东湖校区",
    }


def test_fixed_address_from_config_missing_coords_raises():
    from types import SimpleNamespace

    cfg = SimpleNamespace(
        sss_fixed_lnt=None,
        sss_fixed_lat=None,
        sss_fixed_area_code="330110",
        sss_fixed_address_detail="浙江农林大学东湖校区",
    )
    try:
        sss._fixed_address_from_config(cfg)
    except LookupError:
        pass
    else:
        raise AssertionError("缺少经纬度时应抛出 LookupError")


def test_normalize_timeout_converts_milliseconds():
    from app.api_client import _normalize_timeout

    connect, read = _normalize_timeout(8000)
    assert (connect, read) == (5.0, 8.0)
    assert _normalize_timeout(15) == (5.0, 15.0)
    assert _normalize_timeout((3.0, 9.0)) == (3.0, 9.0)


def test_match_record_prefers_named_fields():
    records = [
        {"id": 1, "remark": "一口轻食"},
        {"id": 2, "name": "一口轻食"},
    ]
    assert sss._match_record(records, "一口轻食", "门店", fields=("name",))["id"] == 2


def test_collect_tasks_reuses_expected_time():
    orders = {"午餐": [
        {"row": 3, "name": "A", "door": "1", "phone": "13800000001"},
        {"row": 4, "name": "B", "door": "2", "phone": "13800000002"},
    ]}
    address = {"lnt": 1.0, "lat": 2.0, "areaCode": "330110", "addressDetail": "X"}
    tasks = sss._collect_tasks(orders, 99, address, "轻食")
    assert len(tasks) == 2
    assert tasks[0]["payload"]["expectedDeliveryTime"] == tasks[1]["payload"]["expectedDeliveryTime"]


def test_submit_tasks_concurrent_faster_than_serial():
    import time

    tasks = [{"identifier": f"t{i}", "payload": {"i": i}} for i in range(10)]

    def submit(payload):
        time.sleep(0.05)
        return {"success": True}

    class Stop:
        def is_set(self):
            return False

    t0 = time.perf_counter()
    result = sss._submit_tasks_concurrent(tasks, submit, Stop(), None, max_workers=5)
    concurrent_elapsed = time.perf_counter() - t0
    assert result.succeeded == {f"t{i}" for i in range(10)} and result.failures == []
    t0 = time.perf_counter()
    result = sss._submit_tasks_concurrent(tasks, submit, Stop(), None, max_workers=1)
    serial_elapsed = time.perf_counter() - t0
    assert result.succeeded == {f"t{i}" for i in range(10)} and result.failures == []
    assert concurrent_elapsed < serial_elapsed / 2


def test_submit_tasks_concurrent_reports_failures():
    tasks = [{"identifier": f"t{i}", "payload": {"i": i}} for i in range(4)]

    def submit(payload):
        if payload["i"] == 2:
            return {"success": False, "message": "boom"}
        return {"success": True}

    class Stop:
        def is_set(self):
            return False

    result = sss._submit_tasks_concurrent(tasks, submit, Stop(), None, max_workers=2)
    assert result.succeeded == {"t0", "t1", "t3"}
    assert result.failures == [("t2", "boom")]


def test_submit_tasks_concurrent_auth_expired_aborts():
    tasks = [{"identifier": f"t{i}", "payload": {"i": i}} for i in range(4)]

    def submit(payload):
        raise sss._AuthExpired("401")

    class Stop:
        def is_set(self):
            return False

    result = sss._submit_tasks_concurrent(tasks, submit, Stop(), None, max_workers=1)
    assert result.auth_error == "401"
    assert result.succeeded == set()


def test_cached_store_id_hit_and_miss():
    from types import SimpleNamespace

    assert sss._cached_store_id(SimpleNamespace(sss_store_id=5, sss_store_name_cached="一口轻食"), "一口轻食") == 5
    assert sss._cached_store_id(SimpleNamespace(sss_store_id=5, sss_store_name_cached="别家"), "一口轻食") is None
    assert sss._cached_store_id(SimpleNamespace(sss_store_id=None, sss_store_name_cached="一口轻食"), "一口轻食") is None


def test_prepare_store_and_address_uses_cache(tmp_path):
    from types import SimpleNamespace

    calls = []

    def fetch_json(path):
        calls.append(path)
        raise AssertionError("命中缓存时不应发起任何查询")

    cfg = SimpleNamespace(sss_store_id=211053, sss_store_name_cached="一口轻食",
                          sss_fixed_lnt=1.0, sss_fixed_lat=2.0,
                          sss_fixed_area_code="330110", sss_fixed_address_detail="X")
    store_id, address = sss._prepare_store_and_address(
        fetch_json, cfg, "一口轻食", "", True, None)
    assert store_id == 211053
    assert address["lnt"] == 1.0
    assert calls == []


def test_dry_run_skips_network(tmp_path):
    from openpyxl import Workbook
    from types import SimpleNamespace

    wb = Workbook()
    ws = wb.active
    ws.title = "午餐"
    ws.append(["姓名", "门牌号", "电话", "送达时间"])
    ws.append(["占位", "", "", ""])
    ws.append(["张三", "A101", "13800000001", "11:00"])
    path = tmp_path / "闪时送.xlsx"
    wb.save(path)
    wb.close()

    cfg = SimpleNamespace(
        sss_excel_path=str(path), sss_account="18758187837",
        sss_dry_run=True, sss_store_name="一口轻食", sss_common_address="嗯哼",
        sss_use_fixed_address=True, sss_fixed_lnt=1.0, sss_fixed_lat=2.0,
        sss_fixed_area_code="330110", sss_fixed_address_detail="X",
        sss_product_name="轻食", api_mode=True, element_timeout_ms=8000,
        sss_url="https://example.invalid", sss_store_id=None,
        sss_store_name_cached="", sss_max_workers=8,
    )

    class Stop:
        def is_set(self):
            return False

    logs = []
    result = sss.run_sss_job(cfg, Stop(), logs.append, password="x",
                             captcha_callback=lambda img: (_ for _ in ()).throw(
                                 AssertionError("干跑不应索取验证码")))
    assert result == {"processed": 1, "created": 1}
    assert any("干跑" in msg for msg in logs)


def test_submit_tasks_concurrent_auth_expired_keeps_successes():
    """401 中断后结构化结果保留已成功集合，主线程可只补提剩余项。"""
    tasks = [{"identifier": f"t{i}", "payload": {"i": i}} for i in range(4)]
    calls = []

    def submit(payload):
        calls.append(payload["i"])
        if payload["i"] == 0:
            return {"success": True}
        raise sss._AuthExpired("401")

    class Stop:
        def is_set(self):
            return False

    result = sss._submit_tasks_concurrent(tasks, submit, Stop(), None, max_workers=1)
    assert result.auth_error == "401"
    assert result.succeeded == {"t0"}


def test_check_success_detects_balance_depleted():
    for msg in ["余额不足，请充值", "账户余额不足", "Balance insufficient", "欠费停服"]:
        try:
            sss._check_success({"success": False, "message": msg})
        except sss._BalanceDepleted:
            pass
        else:
            raise AssertionError(f"应识别为余额不足：{msg}")
    try:
        sss._check_success({"success": False, "message": "地址无效"})
    except sss._BalanceDepleted:
        raise AssertionError("普通失败不应误判为余额不足")
    except LookupError:
        pass
    sss._check_success({"success": True})


def test_submit_tasks_concurrent_balance_aborts_batch():
    tasks = [{"identifier": f"t{i}", "payload": {"i": i}} for i in range(4)]

    def submit(payload):
        if payload["i"] == 1:
            return {"success": False, "message": "余额不足，请充值"}
        return {"success": True}

    class Stop:
        def __init__(self):
            self.flag = False

        def is_set(self):
            return self.flag

        def set(self):
            self.flag = True

    result = sss._submit_tasks_concurrent(tasks, submit, Stop(), None, max_workers=1)
    assert "余额不足" in result.balance_error
    assert result.succeeded == {"t0"}


def test_query_balance_parses_amounts():
    total, frozen = sss.query_balance(
        lambda p: {"success": True, "result": {"totalAmount": 323.9, "freezeAmount": -1189.4}})
    assert (total, frozen) == (323.9, -1189.4)
    assert sss.query_balance(lambda p: {"success": False}) == (None, None)

    def boom(path):
        raise RuntimeError("net down")

    assert sss.query_balance(boom) == (None, None)


def _task(identifier="t1", *, name="张三", phone="13800000001", door="B1",
          delivery="2026-09-10 11:00:00"):
    payload = {
        "receiveName": name,
        "receivePhone": phone,
        "expectedDeliveryTime": delivery,
        "receiveAddress": {"doorNum": door},
    }
    return {"identifier": identifier, "payload": payload,
            "fingerprint": sss._payload_fingerprint(payload)}


def _station_record(task):
    payload = task["payload"]
    return {
        "receiveName": payload["receiveName"],
        "receivePhone": payload["receivePhone"],
        "expectedDeliveryTime": payload["expectedDeliveryTime"],
        "receiveAddress": {"doorNum": payload["receiveAddress"]["doorNum"]},
    }


def test_sss_client_retries_get_but_not_post():
    from app.api_client import SssApiClient

    client = SssApiClient("https://example.invalid/takeout", "a", "p")
    try:
        adapter = client.session.get_adapter("https://example.invalid")
        assert set(adapter.max_retries.allowed_methods) == {"GET"}
    finally:
        client.close()


def test_sss_client_treats_non_json_401_as_auth_expired():
    from app.api_client import ApiError, SssApiClient

    class Html401:
        status_code = 401

        def json(self):
            raise ValueError("HTML login page")

    client = SssApiClient("https://example.invalid/takeout", "a", "p")
    client.session.request = lambda *args, **kwargs: Html401()  # type: ignore[method-assign]
    try:
        with pytest.raises(ApiError, match="401"):
            client.post_json("/consumer/order/one-touch-send/create-order-from-client", {})
    finally:
        client.close()


def test_api_submitter_uses_forked_worker_session_and_closes_it():
    class Worker:
        def __init__(self):
            self.closed = False

        def post_json(self, path, payload):
            assert path == sss._CREATE_ORDER_PATH
            return {"success": True, "payload": payload}

        def close(self):
            self.closed = True

    class Root:
        def __init__(self):
            self.workers = []

        def fork(self):
            worker = Worker()
            self.workers.append(worker)
            return worker

    root = Root()
    submit, close = sss._make_api_submitter(root)
    assert submit({"x": 1})["success"] is True
    assert len(root.workers) == 1
    close()
    assert root.workers[0].closed is True


def test_reconcile_normalizes_door_and_keeps_lunch_dinner_separate():
    lunch = _task("lunch", door="b1", delivery="2026-09-10 11:00:00")
    dinner = _task("dinner", door="B1", delivery="2026-09-10 17:00:00")
    lunch_start, _ = sss._delivery_day_bounds("2026-09-10 11:00:00")

    def fetch(path):
        records = [_station_record(lunch)] if f"startTime={lunch_start}" in path else []
        return {"success": True, "result": {"records": records, "total": len(records)}}

    result = sss._reconcile_tasks([lunch, dinner], fetch)
    assert result.confirmed == {"lunch"}
    assert [task["identifier"] for task in result.missing] == ["dinner"]
    assert result.duplicate_count == 0


def test_list_pending_orders_paginates():
    task = _task()
    calls = []
    unrelated = [{"receiveName": f"u{i}", "receivePhone": "1",
                  "expectedDeliveryTime": "2026-09-10 11:00:00",
                  "receiveAddress": {"doorNum": "X"}} for i in range(100)]

    def fetch(path):
        calls.append(path)
        if "pageNo=1" in path:
            return {"success": True, "result": {"records": unrelated, "total": 101}}
        return {"success": True, "result": {"records": [_station_record(task)], "total": 101}}

    result = sss._reconcile_tasks([task], fetch)
    assert result.confirmed == {"t1"}
    assert len(calls) == 2
    assert all("sortType=1" in path and "sort=1" in path for path in calls)
    assert all("statusList=0" in path and "statusList=1" in path and "statusList=2" in path
               for path in calls)


def test_reconcile_fails_closed_when_list_record_cannot_build_fingerprint():
    task = _task()
    malformed = {
        "receiveName": "张三",
        "receivePhone": "13800000001",
        "expectedDeliveryTime": "2026-09-10 11:00:00",
        "receiveAddress": {},
    }
    with pytest.raises(LookupError, match="无法安全对账"):
        sss._reconcile_tasks(
            [task], lambda path: {"success": True, "result": {"list": [malformed], "total": 1}})


def test_reconcile_detects_extra_duplicate_and_duplicate_excel_rows():
    first = _task("first")
    second = _task("second")
    records = [_station_record(first), _station_record(first), _station_record(first)]

    result = sss._reconcile_tasks(
        [first, second], lambda path: {"success": True, "result": {"records": records, "total": 3}})
    assert result.confirmed == {"first", "second"}
    assert result.missing == []
    assert result.duplicate_count == 1


def test_submit_stop_does_not_schedule_new_tasks_and_closes_sessions():
    calls = []
    closed = []

    class Stop:
        flag = False

        def is_set(self):
            return self.flag

        def set(self):
            self.flag = True

    stop = Stop()

    def submit(payload):
        calls.append(payload["i"])
        stop.set()
        return {"success": True}

    tasks = [{"identifier": f"t{i}", "payload": {"i": i}} for i in range(4)]
    result = sss._submit_tasks_concurrent(tasks, submit, stop, None, max_workers=1,
                                          on_stop=lambda: closed.append(True))
    assert calls == [0]
    assert closed == [True]
    assert result.succeeded == {"t0"}
    assert result.stopped is True


def test_uncertain_post_is_confirmed_by_reconciliation_without_retry():
    task = _task()
    on_site = []
    submits = []

    def fetch(path):
        return {"success": True, "result": {"records": list(on_site), "total": len(on_site)}}

    def factory():
        def submit(payload):
            submits.append(payload)
            on_site.append(_station_record(task))
            raise sss._SubmissionUncertain("ReadTimeout")

        return submit, lambda: None

    class Stop:
        def is_set(self):
            return False

        def set(self):
            raise AssertionError("已对账成功时不应停止")

    final, reconciled = sss._run_reconciled_submission(
        [task], factory, fetch, Stop(), None, None, max_workers=1)
    assert reconciled is True
    assert final is not None and final.confirmed == {"t1"}
    assert len(submits) == 1


def test_auth_expiry_relogs_once_then_only_submits_missing_task():
    task = _task()
    on_site = []
    submits = []
    relogins = []

    def fetch(path):
        return {"success": True, "result": {"records": list(on_site), "total": len(on_site)}}

    def factory():
        def submit(payload):
            submits.append(payload)
            if len(submits) == 1:
                raise sss._AuthExpired("401")
            on_site.append(_station_record(task))
            return {"success": True}

        return submit, lambda: None

    class Stop:
        def is_set(self):
            return False

        def set(self):
            raise AssertionError("401 恢复成功时不应停止")

    final, reconciled = sss._run_reconciled_submission(
        [task], factory, fetch, Stop(), None, None, max_workers=1,
        relogin=lambda: relogins.append(True))
    assert reconciled is True
    assert final is not None and final.confirmed == {"t1"}
    assert len(submits) == 2
    assert relogins == [True]


# ---------------------------------------------------------------------------
# P0 订单幂等性：完整指纹、对账延迟、并发互斥
# ---------------------------------------------------------------------------

def _full_task(identifier="full", *, store_id=211053, goods_name="轻食",
               address_detail="浙江农林大学东湖校区", door="A101",
               account="18758187837", created_at=None):
    payload = {
        "receiveName": "张三",
        "receivePhone": "13800000001",
        "expectedDeliveryTime": "2026-09-10 11:00:00",
        "orderType": 2,
        "storeId": store_id,
        "goodsDetail": [{"goodsName": goods_name, "goodsNum": 1}],
        "receiveAddress": {
            "lnt": 119.728224,
            "lat": 30.256632,
            "areaCode": "330110",
            "addressDetail": address_detail,
            "doorNum": door,
        },
    }
    task = {
        "identifier": identifier,
        "payload": payload,
        "account": account,
        "fingerprint": sss._payload_fingerprint(payload, account=account),
    }
    if created_at is not None:
        task["created_at"] = created_at
    return task


def _full_station_record(task, **overrides):
    payload = task["payload"]
    address = payload["receiveAddress"]
    record = {
        "id": 1,
        "orderSn": "SN-1",
        "receiveName": payload["receiveName"],
        "receivePhone": payload["receivePhone"],
        "expectedDeliveryTime": payload["expectedDeliveryTime"],
        "orderType": payload["orderType"],
        "storeId": payload["storeId"],
        "goodsDetail": payload["goodsDetail"],
        "receiveAddress": dict(address),
        "user": {"mobile": task.get("account", "")},
        "created_at": task.get("created_at") or "2026-09-09 10:00:00",
    }
    record.update(overrides)
    return record


def test_payload_fingerprint_contains_all_context_fields():
    task = _full_task()
    fp = sss._payload_fingerprint(task["payload"], account=task["account"])
    assert fp.store_id == "211053"
    assert fp.goods_name == "轻食"
    assert fp.goods_num == "1"
    assert fp.address_detail == "浙江农林大学东湖校区"
    assert fp.area_code == "330110"
    assert fp.lnt == "119.728224"
    assert fp.lat == "30.256632"
    assert fp.order_type == "2"
    assert fp.account == "18758187837"
    assert sss._order_record_fingerprint(_full_station_record(task)) == fp


@pytest.mark.parametrize("change", [
    {"storeId": 999},
    {"goodsDetail": [{"goodsName": "豪华餐", "goodsNum": 1}]},
    {"receiveAddress": {"lnt": 1.0, "lat": 2.0, "areaCode": "330110",
                        "addressDetail": "别的地址", "doorNum": "A101"}},
    {"receiveAddress": {"lnt": 119.728224, "lat": 30.256632, "areaCode": "999999",
                        "addressDetail": "浙江农林大学东湖校区", "doorNum": "A101"}},
])
def test_reconcile_requires_full_fingerprint(change):
    task = _full_task()
    record = _full_station_record(task)
    record.update(change)
    result = sss._reconcile_tasks(
        [task], lambda path: {"success": True, "result": {"records": [record], "total": 1}})
    assert result.confirmed == set()
    assert [item["identifier"] for item in result.missing] == ["full"]


def test_reconcile_rejects_similar_order_created_before_batch_window():
    task = _full_task()
    record = _full_station_record(task, created_at="2020-01-01 00:00:00")
    result = sss._reconcile_tasks(
        [task], lambda path: {"success": True, "result": {"records": [record], "total": 1}},
        created_after=1_700_000_000.0)
    assert result.confirmed == set()
    assert [item["identifier"] for item in result.missing] == ["full"]


def test_reconcile_delay_polls_without_duplicate_post(monkeypatch):
    task = _full_task()
    on_site = []
    submits = []
    fetches = []

    def fetch(path):
        fetches.append(path)
        # 初始对账（1 次）后，服务端列表延迟两次，第三次才返回订单。
        if len(fetches) < 3:
            return {"success": True, "result": {"records": [], "total": 0}}
        return {"success": True, "result": {"records": [dict(on_site[0])], "total": 1}}

    def factory():
        def submit(payload):
            submits.append(payload)
            on_site.append(_full_station_record(
                task, created_at=int(dt.datetime.now().timestamp() * 1000)))
            return {"success": True}

        return submit, lambda: None

    class Stop:
        def is_set(self):
            return False

        def set(self):
            raise AssertionError("对账延迟不应停止任务")

    monkeypatch.setattr(sss, "_RECONCILE_POLL_INTERVAL_S", 0)
    final, reconciled = sss._run_reconciled_submission(
        [task], factory, fetch, Stop(), None, None, max_workers=1)
    assert reconciled is True
    assert final is not None and final.confirmed == {"full"}
    assert len(submits) == 1, "对账延迟期间禁止重复 POST"
    assert len(fetches) >= 3


def test_sss_job_lock_rejects_concurrent_run(tmp_path):
    from types import SimpleNamespace

    class Stop:
        def is_set(self):
            return False

    sss._SSS_RUN_LOCK.acquire()
    try:
        with pytest.raises(RuntimeError, match="拒绝并发执行"):
            sss.run_sss_job(SimpleNamespace(sss_excel_path=str(tmp_path / "missing.xlsx")),
                            Stop())
    finally:
        sss._SSS_RUN_LOCK.release()


def test_auth_expiry_preserves_successes_confirmed_after_relogin():
    first = _full_task("first")
    second = _full_task("second", door="B202")
    on_site = []
    submits = []
    relogins = []

    def fetch(path):
        return {"success": True, "result": {"records": list(on_site), "total": len(on_site)}}

    def factory():
        def submit(payload):
            submits.append(payload)
            if len(submits) == 1:
                on_site.append(_full_station_record(
                    first, created_at=int(dt.datetime.now().timestamp() * 1000)))
                return {"success": True}
            if len(submits) == 2:
                raise sss._AuthExpired("401")
            on_site.append(_full_station_record(
                second, created_at=int(dt.datetime.now().timestamp() * 1000)))
            return {"success": True}

        return submit, lambda: None

    class Stop:
        def is_set(self):
            return False

        def set(self):
            raise AssertionError("401 恢复成功时不应停止")

    final, reconciled = sss._run_reconciled_submission(
        [first, second], factory, fetch, Stop(), None, None, max_workers=1,
        relogin=lambda: relogins.append(True))
    assert reconciled is True
    assert final is not None and final.confirmed == {"first", "second"}
    assert len(submits) == 3
    assert relogins == [True]


def test_collect_tasks_reuses_stable_client_request_id_on_retry():
    orders = {"午餐": [{"row": 3, "name": "张三", "door": "A101",
                        "phone": "13800000001"}]}
    address = {"lnt": 1.0, "lat": 2.0, "areaCode": "330110", "addressDetail": "X"}
    first = sss._collect_tasks(orders, 99, address, "轻食",
                               account="acct", batch_id="batch-a")
    second = sss._collect_tasks(orders, 99, address, "轻食",
                                account="acct", batch_id="batch-b")
    assert first[0]["client_request_id"] == second[0]["client_request_id"]
    assert "clientRequestId" not in first[0]["payload"]

    enabled = sss._collect_tasks(orders, 99, address, "轻食", account="acct",
                                 idempotency_field="clientRequestId")
    assert enabled[0]["payload"]["clientRequestId"] == enabled[0]["client_request_id"]


# ---------------------------------------------------------------------------
# 闪时送登录态失效识别（HTTP 200 + code=10000）
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("payload", [
    {"success": False, "message": "token失效，请重新登陆", "code": 10000},
    {"success": False, "message": "登录态失效，请重新登录", "code": 10000},
    {"success": False, "message": "TOKEN INVALID", "code": 500},
    {"success": False, "message": "Unauthorized", "code": 500},
])
def test_sss_client_detects_auth_expired_payload(payload):
    from app.api_client import ApiError, SssApiClient

    class Response:
        status_code = 200

        def json(self):
            return payload

    client = SssApiClient("https://example.invalid/takeout", "a", "p")
    client.session.request = lambda *args, **kwargs: Response()  # type: ignore[method-assign]
    try:
        with pytest.raises(ApiError, match="登录态已失效|token|TOKEN|Unauthorized"):
            client.get_json(sss._ORDER_LIST_PATH)
    finally:
        client.close()


def test_post_one_maps_code_10000_api_error_to_auth_expired():
    class Client:
        def post_json(self, path, payload):
            raise sss.ApiError("闪时送登录态已失效（code=10000，token失效，请重新登陆），请重新登录")

    with pytest.raises(sss._AuthExpired, match="10000"):
        sss._post_one(Client(), {})


def test_query_balance_propagates_auth_expired_instead_of_silent_none():
    def fetch_auth(path):
        raise sss.ApiError("闪时送登录态已失效（code=10000，token失效），请重新登录")

    with pytest.raises(sss._AuthExpired):
        sss.query_balance(fetch_auth)

    def fetch_payload(path):
        return {"success": False, "message": "登录态失效，请重新登录", "code": 10000}

    with pytest.raises(sss._AuthExpired):
        sss.query_balance(fetch_payload)


def test_prepare_store_and_address_serial_mode_uses_single_thread():
    import threading
    from types import SimpleNamespace

    threads = []
    responses = {
        sss._STORE_LIST_PATH: {"result": {"records": [{"id": 1, "name": "一口轻食"}]}},
        sss._FREQUENT_ADDR_PATH: {"result": {"records": [{
            "contactName": "嗯哼", "longitude": 1.0, "latitude": 2.0,
            "code": "330110", "position": "X"}]}},
    }

    def fetch_json(path):
        threads.append(threading.current_thread().name)
        return responses[path]

    cfg = SimpleNamespace(sss_store_id=None, sss_store_name_cached="")
    store_id, address = sss._prepare_store_and_address(
        fetch_json, cfg, "一口轻食", "嗯哼", False, None, parallel=False)
    assert store_id == 1
    assert address["areaCode"] == "330110"
    assert len(set(threads)) == 1


def test_with_auth_relogin_retries_once_after_token_expired():
    calls = []

    def operation():
        calls.append("operation")
        if calls.count("operation") == 1:
            raise sss._AuthExpired("token失效，请重新登陆 code=10000")
        return "ok"

    result = sss._with_auth_relogin(operation, lambda: calls.append("relogin"), None, "读取余额")
    assert result == "ok"
    assert calls == ["operation", "relogin", "operation"]
