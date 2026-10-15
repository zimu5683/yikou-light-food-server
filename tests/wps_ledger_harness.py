"""WPS 账本兜底相关测试共用的合成云端与排单表。

**非 ``test_`` 前缀**，pytest 不收集它；只放共用的替身，不放用例（与
``tests/r8s3_location_guard_harness.py`` 同一约定）。

隔离：全部合成数据 —— 内存网格 + 临时目录；不联网、不读真实客户数据、
不写真实 WPS 云表、不碰真实用户目录里的账本/日志。
"""
from __future__ import annotations

import datetime as _dt

from app.wps.models import SheetPlan
from app.wps.sync import CloudOrder, build_plan

#: 合成云表列（0-based，与真实排单表结构一致）。
COL_NAME, COL_ADDR, COL_PHONE = 0, 1, 2
COL_TOTAL, COL_SERVED, COL_LEFT, COL_REMARK = 7, 8, 9, 10
#: 目标日期列（0-based）：表头写「9.11 周五」，目标日期就是 2026-09-11。
COL_TARGET = 4
#: 备注区右侧再放一个同名日期表头：程序只认「电话列 ~ 类型列」之间的日期列，
#: 这个不能被发现（回归 D01 的列定位口径）。
COL_TARGET_DECOY = 12

SHEET = "东湖中餐"
FILE_ID = "F1"
TARGET = _dt.date(2026, 9, 11)
WEEKDAY_NAMES = ("周一", "周二", "周三", "周四", "周五", "周六", "周日")


def date_header(target: _dt.date = TARGET) -> str:
    return f"{target.month}.{target.day} {WEEKDAY_NAMES[target.weekday()]}"


def base_header(target: _dt.date = TARGET) -> dict[int, str]:
    return {
        COL_NAME: "名字", COL_ADDR: "地址", COL_PHONE: "电话",
        3: "9.10 周四", COL_TARGET: date_header(target),
        5: "类型", 6: "餐种", COL_TOTAL: "总餐次",
        COL_SERVED: "已出餐", COL_LEFT: "剩余餐", COL_REMARK: "备注",
        COL_TARGET_DECOY: date_header(target),
    }


def make_grid(header: dict[int, str] | None = None,
              rows: list[dict[int, str]] | None = None,
              target: _dt.date = TARGET) -> dict[tuple[int, int], str]:
    """网格键是 ``(行下标, 列下标)``：``1`` = 表头行（真实第 2 行），``2`` 起是数据。"""
    grid: dict[tuple[int, int], str] = {}
    for col, text in (header if header is not None else base_header(target)).items():
        grid[(1, col)] = text
    for index, row in enumerate(rows or [], start=2):
        for col, text in row.items():
            if text not in (None, ""):
                grid[(index, col)] = str(text)
    return grid


class GridCloud:
    """``KdocsCli`` 的内存替身：只读网格 + 记录写入，不联网、不写真实文档。"""

    path = "/fake/kdocs-cli"

    def __init__(self, grid: dict[tuple[int, int], str]) -> None:
        self.grid = dict(grid)
        self.writes: list[list[dict]] = []
        self.inserts: list[tuple[int, int]] = []

    # -- 只读 --------------------------------------------------------
    def authenticated(self) -> bool:
        return True

    def sheets_info(self, file_id: str):
        return [{"sheetId": 1, "sheetName": "Sheet1", "rowTo": 200, "colTo": 50}]

    def read_grid(self, file_id, worksheet_id, row_from, row_to, col_from, col_to,
                  *, with_format: bool = False):
        hits = {key: value for key, value in self.grid.items()
                if row_from <= key[0] <= row_to and col_from <= key[1] <= col_to}
        if with_format:
            return {key: {"text": str(value), "fill": ""} for key, value in hits.items()}
        return hits

    def read_formulas(self, file_id, worksheet_id, row_from, row_to, col_from, col_to):
        return {key: str(value) for key, value in self.grid.items()
                if row_from <= key[0] <= row_to and col_from <= key[1] <= col_to
                and str(value).startswith("=")}

    def read_cell_format(self, file_id, worksheet_id, row, col):
        return None

    # -- 写入（只改内存） --------------------------------------------
    def write_cells(self, file_id, worksheet_id, cells):
        self.writes.append([dict(cell) for cell in cells])
        for cell in cells:
            self.grid[(int(cell["row"]) - 1, int(cell["col"]) - 1)] = str(cell["value"])

    def insert_rows(self, file_id, worksheet_id, *, row, count):
        self.inserts.append((row, count))
        self.grid = {(r + count if r >= row - 1 else r, c): v
                     for (r, c), v in self.grid.items()}

    def delete_rows(self, file_id, worksheet_id, *, row, count):
        lo, hi = row - 1, row - 1 + count - 1
        shifted = {}
        for (r, c), value in self.grid.items():
            if r < lo:
                shifted[(r, c)] = value
            elif r > hi:
                shifted[(r - count, c)] = value
        self.grid = shifted

    def write_format_ops(self, file_id, worksheet_id, ops):
        return None

    def sort_range(self, file_id, worksheet_id, **kwargs):
        return None

    def delete_columns(self, file_id, worksheet_id, **kwargs):
        return None


def orders(*meals: int, name: str = "张", phone: str = "111",
           address: str = "小", sheet: str = SHEET) -> list[CloudOrder]:
    """同一个人的若干行（一餐一行）；每行算 1 餐，由 planner 相加。"""
    return [CloudOrder(sheet, name, address, phone, "中餐", "经济", meal)
            for meal in meals]


def plan_without_ledger(cli, order_list, *, target: _dt.date = TARGET,
                        sort: bool = False) -> list[SheetPlan]:
    """按**空账本**建计划 —— 正是账本重建要推断的那份输入。"""
    return build_plan(
        cli, local_orders={SHEET: order_list},
        tables={SHEET: {"file_id": FILE_ID}}, target=target, ledger=None,
        address_order={}, sort_enabled=sort)


def write_local_excel(path, *, name: str = "合成客户", phone: str = "13800000000",
                      meals: int = 2, address: str = "西溪北苑",
                      order_no: str = "A-1", target: _dt.date = TARGET) -> None:
    """用真实模板生成排单表再填一行数据；读回走真实 ``read_local_orders``。"""
    from openpyxl import load_workbook

    from app.order.templates import write_order_template

    write_order_template(path)
    workbook = load_workbook(path)
    try:
        sheet = workbook[SHEET]
        sheet.cell(row=3, column=1, value=order_no)
        sheet.cell(row=3, column=2, value=name)
        sheet.cell(row=3, column=3, value=address)
        sheet.cell(row=3, column=4, value=phone)
        # 周一到周日 7 列（第 5~11 列）：勾上"目标日期是周几"，通过批次日期核对。
        sheet.cell(row=3, column=5 + target.weekday(), value="1")
        sheet.cell(row=3, column=12, value="中餐")
        sheet.cell(row=3, column=13, value="经济")
        sheet.cell(row=3, column=14, value=meals)
        workbook.save(path)
    finally:
        workbook.close()
