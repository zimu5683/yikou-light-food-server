"""整理订单故障与恢复回归：清表失败、取消、历史补单、原子保存、分页与部分失败。

全部使用合成订单、临时 Excel 与伪客户端，不触网、不读真实名单/凭据。
"""
from __future__ import annotations

import datetime as dt
import re
import threading
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.order import runner as automod
from app.order import excel_io as excel_io_module
from app.order.excel_io import (
    OrderSaveError,
    _atomic_save_workbook,
    _save_workbook_with_retry,
)
from app.order.fetching import _state_code, split_refund_orders
from app.order.runner import _api_list_waimai_orders, run_job
from app.order.templates import write_order_template

GOOD_ADDRESS = "浙江农林大学东湖校区 A5 506"


def _prepare_excel(tmp_path: Path) -> Path:
    excel = tmp_path / "排单.xlsx"
    write_order_template(excel)
    wb = load_workbook(excel)
    ws = wb["东湖中餐"]
    ws.cell(3, 1).value = "OLD"
    ws.cell(3, 2).value = "旧姓名"
    ws.cell(3, 3).value = "原数据"
    ws.cell(3, 4).value = "13900000000"
    wb.save(excel)
    wb.close()
    return excel


def _config(excel: Path, date_text: str):
    config = type("FakeConfig", (), {})()
    config.excel_path = str(excel)
    config.target_url = "https://example.invalid"
    config.phone_number = "13900000000"
    config.order_date = date_text
    return config


def _install_temp_aliases(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(automod, "load_aliases", lambda: {})
    monkeypatch.setattr(automod, "aliases_path", lambda: tmp_path / "aliases.json")


def _fake_client(api_get):
    class FakeClient:
        def __init__(self, *args, **kwargs):
            pass

        def login(self):
            pass

        def get_json(self, path):
            return api_get(path)

    return FakeClient


def _order_list(entries):
    """构造一页订单列表响应；entries 为 (order_id, pick_no, date_text, state)。"""
    return {"data": {"list": [
        {"id": oid, "pickNo": pick, "storeId": "1",
         "created_at": f"{date_text} 10:00:00", "state": state}
        for oid, pick, date_text, state in entries
    ], "total": len(entries)}}


def _detail(order_id, code, address=GOOD_ADDRESS, goods="单点经济餐（午餐）"):
    return {"data": {
        "id": order_id,
        "address": {"contact": f"顾客{code}", "mobile": "13800000000",
                    "address": address},
        "goods": [{"name": goods, "num": 1}],
    }}


def test_clear_failure_aborts_without_saving_polluted_workbook(tmp_path, monkeypatch):
    """清旧表失败时必须整体中止：不清空、不追加、不保存污染结果。"""
    excel = _prepare_excel(tmp_path)
    today = dt.date.today()
    date_text = today.isoformat()

    def api_get(path):
        if "/channel/order?" in path and "pageNo=1" in path:
            return _order_list([("103", "W6", date_text, 6)])
        if "/channel/order/103" in path:
            return _detail("103", "W6")
        return {"data": {}}

    monkeypatch.setattr(automod, "AdminApiClient", _fake_client(api_get))
    _install_temp_aliases(monkeypatch, tmp_path)

    def boom(*args, **kwargs):
        raise RuntimeError("清表模拟失败")

    monkeypatch.setattr(automod, "clear_campus_sub_sheets", boom)
    saves: list[bool] = []
    monkeypatch.setattr(automod, "_save_workbook_with_retry",
                        lambda *args, **kwargs: saves.append(True))
    monkeypatch.setattr(automod, "write_pending", lambda *args, **kwargs: (
        (_ for _ in ()).throw(AssertionError("中止时不应写待确认报告"))))

    result = run_job(_config(excel, date_text), None, threading.Event(), [],
                     password="pw", save_decision_callback=lambda _: "cancel")

    assert result["status"] == "failed"
    assert result["abort_reason"] == "clear_failed"
    assert result["summary"]["confirmed"] == 0
    assert result["summary"]["failed"] == 1
    assert "原文件未被修改" in result["next_action"]
    assert saves == []
    backups = list((tmp_path / "backups").glob("排单_*.xlsx"))
    assert backups, "清表失败中止前也应保留原文件备份"
    wb = load_workbook(excel)
    ws = wb["东湖中餐"]
    assert [ws["A3"].value, ws["C3"].value] == ["OLD", "原数据"]
    assert ws["A4"].value is None
    wb.close()


def test_cancel_during_detail_preserves_original_campus_sheet(tmp_path, monkeypatch):
    """取消/停止后不能清空原表，也不能保存本轮修改。"""
    excel = _prepare_excel(tmp_path)
    today = dt.date.today()
    date_text = today.isoformat()
    stop_event = threading.Event()
    clear_calls: list[bool] = []
    save_calls: list[bool] = []

    def api_get(path):
        if "/channel/order?" in path and "pageNo=1" in path:
            return _order_list([("103", "W6", date_text, 6)])
        if "/channel/order/103" in path:
            stop_event.set()  # 用户在详情阶段取消
            return _detail("103", "W6")
        return {"data": {}}

    monkeypatch.setattr(automod, "AdminApiClient", _fake_client(api_get))
    _install_temp_aliases(monkeypatch, tmp_path)
    monkeypatch.setattr(automod, "write_pending", lambda *args, **kwargs: tmp_path / "pending.json")
    monkeypatch.setattr(automod, "clear_campus_sub_sheets",
                        lambda wb: clear_calls.append(True) or [])
    monkeypatch.setattr(automod, "_save_workbook_with_retry",
                        lambda *args, **kwargs: save_calls.append(True))

    result = run_job(_config(excel, date_text), None, stop_event, [],
                     password="pw", save_decision_callback=lambda _: "cancel")

    assert result["status"] == "stopped"
    assert result["summary"]["stopped"] is True
    assert clear_calls == []
    assert save_calls == []
    wb = load_workbook(excel)
    assert wb["东湖中餐"]["A3"].value == "OLD"
    wb.close()


def test_historical_target_only_appends_history_keeps_current_campus(tmp_path, monkeypatch):
    """历史补单只追加日期表，绝不能顺带清空当天六张校区表。"""
    excel = _prepare_excel(tmp_path)
    today = dt.date.today()
    target = today - dt.timedelta(days=10)
    target_text = target.isoformat()

    def api_get(path):
        if "/channel/order?" in path and "pageNo=1" in path:
            return _order_list([("103", "W6", target_text, 6)])
        if "/channel/order/103" in path:
            return _detail("103", "W6")
        return {"data": {}}

    monkeypatch.setattr(automod, "AdminApiClient", _fake_client(api_get))
    _install_temp_aliases(monkeypatch, tmp_path)
    monkeypatch.setattr(automod, "write_pending", lambda *args, **kwargs: tmp_path / "pending.json")

    result = run_job(_config(excel, target_text), None, threading.Event(), [],
                     password="pw", save_decision_callback=lambda _: "cancel")

    assert result["status"] == "confirmed"
    assert result["found"] == 1
    wb = load_workbook(excel)
    assert wb["东湖中餐"]["A3"].value == "OLD"
    history_name = (f"{target.year}年{target.month}月{target.day}日 "
                    f"{automod.WEEKDAYS[target.weekday()]}")
    assert history_name in wb.sheetnames
    assert wb[history_name]["A2"].value == "W6"
    assert wb[history_name]["C2"].value == "A5"
    weekday_sheet = wb[automod.WEEKDAYS[(today.weekday() + 1) % 7]]
    assert weekday_sheet["A3"].value is None
    wb.close()


def test_atomic_save_cancel_preserves_original_and_cleans_temp(tmp_path):
    target = tmp_path / "排单.xlsx"
    Workbook().save(target)
    before = target.read_bytes()
    workbook = load_workbook(target)

    def locked(path):
        raise PermissionError("file is locked")

    workbook.save = locked  # type: ignore[method-assign]
    with pytest.raises(PermissionError, match="原文件未被修改"):
        _save_workbook_with_retry(workbook, target, lambda error: "cancel")
    workbook.close()

    assert target.read_bytes() == before
    leftovers = [p.name for p in tmp_path.iterdir() if p.name.startswith(".排单.")]
    assert leftovers == []


def test_atomic_save_retry_writes_temp_then_replaces_target(tmp_path):
    target = tmp_path / "排单.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "before"
    workbook.save(target)
    loaded = load_workbook(target)
    calls: list[Path] = []
    real_save = loaded.save

    def flaky(path):
        calls.append(Path(path))
        if len(calls) == 1:
            raise PermissionError("first save locked")
        real_save(path)

    loaded.save = flaky  # type: ignore[method-assign]
    _save_workbook_with_retry(loaded, target, lambda error: "retry")
    loaded.close()

    assert len(calls) == 2
    assert calls[0] != target and calls[0].suffix == target.suffix
    result = load_workbook(target)
    assert result.active["A1"].value == "before"
    result.close()
    assert [p.name for p in tmp_path.iterdir() if p.name.startswith(".排单.")] == []


def test_split_refund_orders_tolerates_non_numeric_state():
    rows_by_pick = {
        "W8": [{"state": "state=8（已退款）"}],
        "W7": [{"state": "7 pending"}],
        "W6": [{"state": "unknown"}],
        "W5": [{}],
    }
    normal, refunded, applied = split_refund_orders(rows_by_pick, [8, 7, 6, 5])
    assert normal == [6, 5]
    assert refunded == [8]
    assert applied == [7]
    assert _state_code("state=8") == 8
    assert _state_code("unknown") is None


def test_pagination_duplicate_deduped_but_same_person_multiple_orders_kept():
    """重复分页按 order_id 去重；同一人的两个合法 order_id 必须都保留。"""
    target = dt.date(2026, 9, 7)

    def api_get(path):
        page = int(re.search(r"pageNo=(\d+)", path).group(1))
        pages = {
            1: [{"id": "1", "pickNo": "W1", "storeId": "1",
                 "created_at": "2026-09-07 10:00:00"}],
            2: [{"id": "1", "pickNo": "W1", "storeId": "1",
                 "created_at": "2026-09-07 10:00:00"},
                {"id": "2", "pickNo": "W2", "storeId": "1",
                 "created_at": "2026-09-07 10:00:00"}],
        }
        return {"data": {"list": pages.get(page, []), "total": 2}}

    rows = _api_list_waimai_orders(api_get, target, page_size=1, concurrent_pages=True)
    assert [row["order_id"] for row in rows] == ["1", "2"]


def test_run_job_partial_fetch_failure_writes_success_and_reports_partial(tmp_path, monkeypatch):
    excel = _prepare_excel(tmp_path)
    today = dt.date.today()
    date_text = today.isoformat()

    def api_get(path):
        if "/channel/order?" in path and "pageNo=1" in path:
            return _order_list([("1", "W1", date_text, 6), ("2", "W2", date_text, 6)])
        if "/channel/order/2" in path:
            return _detail("2", "W2")
        if "/channel/order/1" in path:
            return {"data": {}}
        return {"data": {}}

    monkeypatch.setattr(automod, "AdminApiClient", _fake_client(api_get))
    _install_temp_aliases(monkeypatch, tmp_path)
    monkeypatch.setattr(automod, "write_pending", lambda *args, **kwargs: tmp_path / "pending.json")
    monkeypatch.setattr(automod.time, "sleep", lambda seconds: None)

    result = run_job(
        _config(excel, date_text), None, threading.Event(), [], password="pw",
        order_decision_callback=lambda code, error: "skip",
        save_decision_callback=lambda error: "cancel",
    )

    assert result["status"] == "partial"
    assert result["found"] == 1
    assert result["failed_orders"] == ["W1"]
    assert result["summary"]["confirmed"] == 1
    assert result["summary"]["unconfirmed"] == 1
    wb = load_workbook(excel)
    values = [cell.value for ws in wb.worksheets for row in ws.iter_rows() for cell in row]
    wb.close()
    assert "W2" in values and "W1" not in values


def test_run_job_all_details_failed_preserves_original(tmp_path, monkeypatch):
    excel = _prepare_excel(tmp_path)
    today = dt.date.today()
    date_text = today.isoformat()
    saves: list[bool] = []

    def api_get(path):
        if "/channel/order?" in path and "pageNo=1" in path:
            return _order_list([("1", "W1", date_text, 6), ("2", "W2", date_text, 6)])
        if "/channel/order/" in path:
            return {"data": {}}
        return {"data": {}}

    monkeypatch.setattr(automod, "AdminApiClient", _fake_client(api_get))
    _install_temp_aliases(monkeypatch, tmp_path)
    monkeypatch.setattr(automod, "write_pending", lambda *args, **kwargs: tmp_path / "pending.json")
    monkeypatch.setattr(automod, "_save_workbook_with_retry",
                        lambda *args, **kwargs: saves.append(True))
    monkeypatch.setattr(automod.time, "sleep", lambda seconds: None)

    result = run_job(
        _config(excel, date_text), None, threading.Event(), [], password="pw",
        order_decision_callback=lambda code, error: "skip",
        save_decision_callback=lambda error: "cancel",
    )

    assert result["status"] == "failed"
    assert result["abort_reason"] == "fetch_failed"
    assert saves == []
    wb = load_workbook(excel)
    assert wb["东湖中餐"]["A3"].value == "OLD"
    wb.close()


def _valid_target(tmp_path: Path) -> tuple[Path, bytes]:
    target = tmp_path / "排单.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "原内容"
    workbook.save(target)
    workbook.close()
    return target, target.read_bytes()


class _NoOutputWorkbook:
    def save(self, path):
        return None  # 不写任何字节也不抛异常


class _EmptyOutputWorkbook:
    def save(self, path):
        Path(path).write_bytes(b"")


class _CorruptOutputWorkbook:
    def save(self, path):
        Path(path).write_bytes(b"not a valid xlsx")


class _RaisingWorkbook:
    def save(self, path):
        raise RuntimeError("disk full")


def test_atomic_save_rejects_no_output_and_preserves_target(tmp_path):
    target, before = _valid_target(tmp_path)
    with pytest.raises(OrderSaveError, match="空文件"):
        _atomic_save_workbook(_NoOutputWorkbook(), target)
    assert target.read_bytes() == before


def test_atomic_save_rejects_empty_output_and_preserves_target(tmp_path):
    target, before = _valid_target(tmp_path)
    with pytest.raises(OrderSaveError, match="空文件"):
        _atomic_save_workbook(_EmptyOutputWorkbook(), target)
    assert target.read_bytes() == before


def test_atomic_save_rejects_corrupt_output_and_preserves_target(tmp_path):
    target, before = _valid_target(tmp_path)
    with pytest.raises(OrderSaveError, match="无法读取/已损坏"):
        _atomic_save_workbook(_CorruptOutputWorkbook(), target)
    assert target.read_bytes() == before


def test_atomic_save_wraps_write_exception_and_preserves_target(tmp_path):
    target, before = _valid_target(tmp_path)
    with pytest.raises(OrderSaveError, match="写临时 Excel 失败"):
        _atomic_save_workbook(_RaisingWorkbook(), target)
    assert target.read_bytes() == before


def test_atomic_save_wraps_temp_creation_failure_and_preserves_target(tmp_path, monkeypatch):
    target, before = _valid_target(tmp_path)

    def denied(*args, **kwargs):
        raise PermissionError("directory not writable")

    monkeypatch.setattr(excel_io_module.tempfile, "mkstemp", denied)
    with pytest.raises(OrderSaveError, match="目录不可写"):
        _atomic_save_workbook(_NoOutputWorkbook(), target)
    assert target.read_bytes() == before


def test_atomic_save_wraps_replace_failure_and_preserves_target(tmp_path, monkeypatch):
    target, before = _valid_target(tmp_path)

    class ValidWorkbook:
        def save(self, path):
            workbook = Workbook()
            workbook.active["A1"] = "新内容"
            workbook.save(path)
            workbook.close()

    def broken_replace(src, dst):
        raise OSError("replace failed")

    monkeypatch.setattr(excel_io_module.os, "replace", broken_replace)
    with pytest.raises(OrderSaveError, match="替换原 Excel 失败"):
        _atomic_save_workbook(ValidWorkbook(), target)
    assert target.read_bytes() == before
    assert [item.name for item in tmp_path.iterdir() if item.name.startswith(".排单.")] == []


def test_atomic_save_replace_permission_error_cancel_preserves_target(tmp_path, monkeypatch):
    target, before = _valid_target(tmp_path)
    calls: list[str] = []

    def locked_replace(src, dst):
        calls.append(str(dst))
        raise PermissionError("target locked")

    class ValidWorkbook:
        def save(self, path):
            workbook = Workbook()
            workbook.active["A1"] = "新内容"
            workbook.save(path)
            workbook.close()

    monkeypatch.setattr(excel_io_module.os, "replace", locked_replace)
    with pytest.raises(PermissionError, match="原文件未被修改"):
        _save_workbook_with_retry(ValidWorkbook(), target, lambda error: "cancel")
    assert calls == [str(target)]
    assert target.read_bytes() == before
