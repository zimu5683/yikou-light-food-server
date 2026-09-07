"""Tests for the Excel template generators (app.excel_templates)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.excel_templates import write_order_template, write_sss_template

ORDER_SHEETS = (
    "东湖中餐", "衣锦中餐", "医学院中餐",
    "东湖晚餐", "衣锦晚餐", "医学院晚餐",
    "周一", "周二", "周三", "周四", "周五", "周六", "周日",
)
SUB_HEADERS = ("订单", "姓名", "地址", "电话", "周一", "周二", "周三",
               "周四", "周五", "周六", "周日", "类型", "餐种", "餐次")
DAY_HEADERS = ("订单", "姓名", "地址", "电话", "餐种", "餐次")


def test_write_order_template_structure(tmp_path: Path):
    dest = tmp_path / "排单.xlsx"
    write_order_template(dest)
    wb = load_workbook(dest)
    assert wb.sheetnames == list(ORDER_SHEETS)
    for name in ORDER_SHEETS[:6]:
        ws = wb[name]
        # 表头在第 2 行、共 14 列；标题行跨 A1:N1 合并。
        assert [ws.cell(2, c).value for c in range(1, 15)] == list(SUB_HEADERS)
        assert "A1:N1" in [str(m) for m in ws.merged_cells.ranges]
    for name in ORDER_SHEETS[6:]:
        ws = wb[name]
        # 周表三区块标题 + 第 2 行两组 6 列表头。
        assert ws["A1"].value == "中餐"
        assert ws["G1"].value == "晚餐"
        assert ws["N1"].value == "总餐"
        for start in (1, 7):
            assert [ws.cell(2, c).value for c in range(start, start + 6)] == list(DAY_HEADERS)
    wb.close()


def test_write_sss_template_structure(tmp_path: Path):
    dest = tmp_path / "闪时送.xlsx"
    write_sss_template(dest)
    wb = load_workbook(dest)
    assert wb.sheetnames == ["午餐", "晚餐"]
    for name in ("午餐", "晚餐"):
        ws = wb[name]
        assert ws["A1"].value == name
        assert [ws.cell(2, c).value for c in range(1, 4)] == ["姓名", "地址", "电话"]
        assert "A1:C1" in [str(m) for m in ws.merged_cells.ranges]
    wb.close()
