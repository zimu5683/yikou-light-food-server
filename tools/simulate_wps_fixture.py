#!/usr/bin/env python3
"""离线仿真：用《测试东湖中餐.xlsx》夹具在本地完整跑一遍云同步新流程。

夹具四张表：
- 东湖中餐（试验田）    = 云端表同步前状态（102 人）→ 程序的写入对象
- 东湖中餐（本地工作表）= 本地排单表（14 个订单）  → 程序的输入
- 东湖中餐（目标）      = 协作者手工维护的旧结果（仅供人工参考，新流程不再比对它：
  新流程会按地址顺序重排**整张表**，行序与协作者手工排的不一定相同）
- 东湖中餐（基准）      = 基线存档

仿真不联网、不改夹具文件：把试验田拷进内存工作簿当"云端"，跑
build_plan + apply_plan，然后按**需求逐条验收**：

1. 计划里每个人都在表里；**总餐次 = 云端原值 + 本地餐次**（累加，不是绝对值）；
   **日期格**：协作者已经写了 0（当天不送）的必须原样保留，其余写 1；
2. 列B 的顺序符合规定顺序（清单外的排最后）；
3. 新行底色 = 模板行底色（经济餐）/ 整行金黄（豪华餐）；
4. 新行公式 = SUM(日期区间)，不含协作者的标记列；
5. 协作者的标记格（第 2 行第 17 列）原样不动；我们的记号写在第 18 列；
6. 行数 = 原人数 + 新客户数；排序辅助列已删除。
7. 同一批**第二次上传零改动**（账本按行记住"本批已同步的本地餐次"）。
8. 同一个人一天下两单（本地两行）时，云端占两行、各标当天 1（这天送两餐）。

用法：.venv/bin/python tools/simulate_wps_fixture.py [夹具路径]
"""
from __future__ import annotations

import datetime as dt
import re
import sys
import tempfile
from copy import copy
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from openpyxl.styles import PatternFill

from app.core.config import default_wps_address_order
from app.wps import sync as wc
from app.wps.sync import CloudOrder, SyncLedger, apply_plan, build_plan, person_key

FIXTURE_DEFAULT = Path("/home/zimu/文档/测试东湖中餐.xlsx")
SHEET_TEST = "东湖中餐（试验田）"
SHEET_LOCAL = "东湖中餐（本地工作表）"

TARGET_DATE = dt.date(2026, 9, 14)      # 夹具表头里有 9.14 周一
RUN_DATE = dt.date(2026, 9, 13)         # 周日运行 → 记号 1
MARK_COL = 17                           # 协作者的日期标记列
MARKER_COL = 18                         # 我们的通讯记号列
GOLD = "FFFFC000"


def cell_fill_hex(cell) -> str:
    """返回单元格底色的 ARGB hex（带 #），无底色返回 ''。"""
    fill = cell.fill
    if fill is None or fill.patternType is None:
        return ""
    rgb = getattr(fill.fgColor, "rgb", None)
    if not isinstance(rgb, str) or len(rgb) != 8:
        return ""
    return f"#{rgb}"


def col_index(letter: str) -> int:
    value = 0
    for char in str(letter).upper():
        value = value * 26 + (ord(char) - 64)
    return value                      # 1-based 列号


def content_last_col(ws) -> int:
    """**有内容**的最后一列（1-based）；空表返回 0。

    不能用 ``ws.max_column``：这个夹具云的 ``read_grid`` 用 ``ws.cell()`` 读格子，
    openpyxl 会顺手把空格子也创建出来 —— 探测排序区时读过一次右侧 40 列，
    ``max_column`` 就虚涨到 58，会把"辅助列没删干净"误报成失败。
    """
    return max((col for (_row, col), cell in ws._cells.items()
                if cell.value not in (None, "")), default=0)


class FixtureCloud:
    """用 openpyxl 工作表模拟云端表的 KdocsCli 替身。

    read_grid 的行为对齐真实接口：只返回非空格；with_format 时附带底色。
    sort_range 按辅助列稳定排序，并把**值与样式一起搬动**（真实接口承诺
    "保留格式与公式"）。delete_columns 左移删除（收尾清辅助列）。
    """

    def __init__(self, ws: openpyxl.worksheet.worksheet.Worksheet):
        self.ws = ws
        self.inserts: list[tuple[int, int]] = []
        self.sorts: list[dict] = []
        self.deleted_columns: list[int] = []
        self.format_ops: list[list[dict]] = []
        self.path = "/fixture/kdocs-cli"

    # ---- KdocsCli 接口 ----

    def sheets_info(self, file_id: str):
        return [{"sheetId": 1, "sheetName": self.ws.title,
                 "rowTo": self.ws.max_row - 1, "colTo": self.ws.max_column - 1}]

    def read_grid(self, file_id, worksheet_id, row_from, row_to, col_from, col_to,
                  *, with_format: bool = False):
        out: dict[tuple[int, int], object] = {}
        for r in range(row_from, row_to + 1):
            for c in range(col_from, col_to + 1):
                cell = self.ws.cell(row=r + 1, column=c + 1)
                if cell.value is None or str(cell.value).strip() == "":
                    continue
                if with_format:
                    out[(r, c)] = {"text": str(cell.value),
                                   "fill": cell_fill_hex(cell)}
                else:
                    out[(r, c)] = str(cell.value)
        return out

    def write_cells(self, file_id, worksheet_id, cells):
        for cell in cells:
            self.ws.cell(row=int(cell["row"]), column=int(cell["col"]),
                         value=str(cell["value"]))

    def read_formulas(self, file_id, worksheet_id, row_from, row_to, col_from, col_to):
        result = {}
        for r in range(row_from, row_to + 1):
            for c in range(col_from, col_to + 1):
                value = self.ws.cell(row=r + 1, column=c + 1).value
                if isinstance(value, str) and value.startswith("="):
                    result[(r, c)] = value
        return result

    def insert_rows(self, file_id, worksheet_id, *, row, count):
        self.inserts.append((row, count))
        self.ws.insert_rows(row, count)          # 在 1-based row 前插 count 行

    def delete_rows(self, file_id, worksheet_id, *, row, count):
        self.ws.delete_rows(row, count)

    def sort_range(self, file_id, worksheet_id, *, range_ref, key, order="asc",
                   header=False, key2=None, order2=None):
        """按辅助列稳定排序区域，值与样式一起搬。"""
        match = re.match(r"^([A-Z]+)(\d+):([A-Z]+)(\d+)$", str(range_ref))
        assert match, f"排序区域格式非法：{range_ref}"
        c0, c1 = col_index(match.group(1)), col_index(match.group(3))
        r0, r1 = int(match.group(2)), int(match.group(4))
        first = r0 + (1 if header else 0)
        key_col = col_index(key)
        self.sorts.append({"range": range_ref, "key": key, "order": order,
                           "header": header})

        snapshot: list[tuple[int, list[tuple[object, object]]]] = []
        for r in range(first, r1 + 1):
            row = []
            for c in range(c0, c1 + 1):
                cell = self.ws.cell(r, c)
                row.append((cell.value, copy(cell._style)))
            snapshot.append((r, row))

        def sort_key(item):
            raw = str(self.ws.cell(item[0], key_col).value or "").strip()
            return (1, "") if not raw else (0, raw)

        ordered = sorted(snapshot, key=sort_key,
                         reverse=str(order).lower() == "desc")
        for offset, (_src, values) in enumerate(ordered):
            target_row = first + offset
            for index, (value, style) in enumerate(values):
                cell = self.ws.cell(target_row, c0 + index)
                cell.value = value
                cell._style = copy(style)

    def delete_columns(self, file_id, worksheet_id, *, column, rows):
        self.deleted_columns.append(column)
        self.ws.delete_cols(column, 1)

    def write_format_ops(self, file_id, worksheet_id, ops):
        self.format_ops.append([dict(op) for op in ops])
        for op in ops:
            color = op["xf"]["fill"]["back"]["value"]
            argb = f"{color & 0xFFFFFFFF:08X}"
            fill = PatternFill("solid", fgColor=argb)
            for r in range(op["rowFrom"] + 1, op["rowTo"] + 2):
                for c in range(op["colFrom"] + 1, op["colTo"] + 2):
                    self.ws.cell(row=r, column=c).fill = fill

    def read_cell_format(self, file_id, worksheet_id, row, col):
        cell = self.ws.cell(row=row, column=col)
        if cell.value is None:
            return None
        font = cell.font
        color = None
        if font.color is not None and getattr(font.color, "type", "") == "rgb":
            color = font.color.rgb
        return {
            "cellText": str(cell.value),
            "fonts": {"font_east_asia": font.name or "", "size": int(font.size or 10),
                      "color": color},
            "alignment": {"horizontal": "haCenter", "vertical": "vaCenter"},
            "cell_background_color": cell_fill_hex(cell),
        }

    def authenticated(self) -> bool:
        return True


def build_local_workbook(fixture: Path, tmp_xlsx: Path) -> None:
    """把夹具的「本地工作表」拷成标准排单表（子表名「东湖中餐」）。"""
    src = openpyxl.load_workbook(fixture, data_only=True)
    ws = src[SHEET_LOCAL]
    dst = openpyxl.Workbook()
    out = dst.active
    out.title = "东湖中餐"
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=14, values_only=True):
        out.append(list(row))
    dst.save(tmp_xlsx)


def norm(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    try:
        return str(int(float(text)))
    except (TypeError, ValueError):
        return text


def is_white(hexfill: str) -> bool:
    return hexfill in ("", "#FFFFFFFF", "#00FFFFFF", "#00000000")


def main() -> int:
    fixture = Path(sys.argv[1]) if len(sys.argv) > 1 else FIXTURE_DEFAULT
    print(f"夹具：{fixture}")
    tmp_xlsx = Path("/tmp/fixture_本地工作表.xlsx")
    build_local_workbook(fixture, tmp_xlsx)

    src = openpyxl.load_workbook(fixture)
    cloud = FixtureCloud(src[SHEET_TEST])

    orders = wc.read_local_orders(tmp_xlsx, sheets=["东湖中餐"])["东湖中餐"]
    # 额外场景（验收第 8 条）：同一个人一天下两单 → 本地两行、云端两行、各标当天 1。
    # 用一个真实本地订单复制出"第二单"（模拟 2026-09-18 那批的雷丝淇/柟）。
    twin_source = orders[0] if orders else None
    if twin_source is not None:
        orders.append(CloudOrder(
            sheet=twin_source.sheet, name=twin_source.name, address=twin_source.address,
            phone=twin_source.phone, meal_type=twin_source.meal_type,
            meal_kind=twin_source.meal_kind, meals=1, row=twin_source.row + 1000,
            order_no=f"{twin_source.order_no}+第二单", rows=(twin_source.row + 1000,),
            weekday_marks=twin_source.weekday_marks))
        print(f"额外场景：{twin_source.name} 同一天两单（第 1 单 {twin_source.meals} 餐 + 第 2 单 1 餐）")
    people_before = [cloud.ws.cell(r, 1).value for r in range(3, cloud.ws.max_row + 1)]
    people_before = [str(n) for n in people_before if n]
    columns_before = content_last_col(cloud.ws)
    mark_before = cloud.ws.cell(2, MARK_COL).value
    totals_before: dict[tuple[str, str], int] = {}
    marks_before: dict[tuple[str, str], str] = {}
    print(f"云端（试验田）：{len(people_before)} 人；本地订单：{len(orders)} 个")

    order_list = default_wps_address_order()["东湖中餐"]
    ledger = SyncLedger(Path(tempfile.mkdtemp()) / "sim_state.json")
    plans = build_plan(
        cloud, local_orders={"东湖中餐": orders},
        tables={"东湖中餐": {"file_id": "FIXTURE", "drive_id": ""}},
        target=TARGET_DATE, ledger=ledger, marker_enabled=True, run_date=RUN_DATE,
        address_order={"东湖中餐": order_list}, sort_enabled=True,
        log=lambda m: print(f"  [log] {m}"))
    plan = plans[0]
    print(f"\n计划：目标列 {plan.target_col}（{plan.target_header}）、"
          f"记号列 {plan.marker_col}→{plan.weekday_number}、"
          f"日期列 {plan.date_cols}、排序区域 {plan.sort_range}")
    for warning in plan.warnings:
        print(f"  ⚠ {warning}")

    template_row = plan.format_rows[0] if plan.format_rows else 0
    template_fills = {c: cell_fill_hex(cloud.ws.cell(template_row, c))
                      for c in range(1, 16)} if template_row else {}

    # 写之前的云端快照（此刻还没写，计划是只读的）：验收要用它当基准。
    # 排序会把行搬走，所以按（姓名, 电话）记。
    for r in range(3, cloud.ws.max_row + 1):
        name = cloud.ws.cell(r, 1).value
        if not name:
            continue
        key = person_key(name, cloud.ws.cell(r, 3).value)
        totals_before[key] = int(float(cloud.ws.cell(r, plan.columns["total"]).value or 0))
        if plan.target_col:
            marks_before[key] = norm(cloud.ws.cell(r, plan.target_col).value)
    kept_zero = sum(1 for value in marks_before.values() if value == "0")

    result = apply_plan(cloud, plans, ledger=ledger, marker_enabled=True,
                        log=lambda m: print(f"  [log] {m}"))
    print(f"\napply 结果：{result['sheets']}")
    if result["failed"]:
        print("❌ 有表写入失败，中止验收")
        return 1

    problems: list[str] = []
    new_changes = [c for c in plan.changes if c.kind == "new"]
    # 索引键必须是（姓名, 电话）——表里有重名的人（「李」「陈」都出现多次）；
    # 一个人可能有多行（一天两单），所以记**行号列表**，第 i 行 = 第 i 个槽位。
    by_person: dict[tuple[str, str], list[int]] = {}
    for r in range(3, cloud.ws.max_row + 1):
        name = cloud.ws.cell(r, 1).value
        if name:
            rows = by_person.setdefault(
                person_key(name, cloud.ws.cell(r, 3).value), [])
            if r not in rows:
                rows.append(r)
    for rows in by_person.values():
        rows.sort()

    def row_of(key, slot: int) -> int:
        rows = by_person.get(key) or []
        return rows[slot - 1] if len(rows) >= slot else 0

    # 1) 计划里的每个人都在表里；总餐次 = 云端原值 + 本次增量；日期格按约定
    for change in plan.changes:
        key = person_key(change.name, change.phone)
        row = row_of(key, change.slot)
        if not row:
            problems.append(f"{change.name} 不在表里")
            continue
        got_mark = norm(cloud.ws.cell(row, plan.target_col).value)
        got_total = norm(cloud.ws.cell(row, plan.columns["total"]).value)
        before_mark = marks_before.get(key, "")
        if change.kind == "new":
            want_mark = wc.CELL_MARK          # 新建的行，日期格是我们写的 1
        else:
            # 老行：协作者已经写了别的值（0 = 当天不送）就保持原样，其余写 1
            want_mark = before_mark if before_mark not in ("", wc.CELL_MARK) else wc.CELL_MARK
        if got_mark != want_mark:
            problems.append(f"{change.name}：日期格应为 {want_mark!r}，实际 {got_mark!r}")
        want_total = totals_before.get(key, 0) + change.local_meals
        if change.kind == "existing" and got_total != str(want_total):
            problems.append(f"{change.name}：总餐次应为 云端 {totals_before.get(key, 0)} "
                            f"+ 本地 {change.local_meals} = {want_total}，实际 {got_total!r}")
        if change.kind == "new" and got_total != str(change.total_after):
            problems.append(f"{change.name}：新客户总餐次应为 {change.total_after}，"
                            f"实际 {got_total!r}")

    # 2) 列B 顺序符合规定顺序（清单外的排最后）
    ranks, tail = wc.build_address_ranks(order_list, [])
    seen_tail = False
    sequence = []
    for r in range(3, cloud.ws.max_row + 1):
        if not cloud.ws.cell(r, 1).value:
            continue
        addr = str(cloud.ws.cell(r, 2).value or "").strip()
        rank = ranks.get(wc._address_key(addr), tail)
        sequence.append((r, addr, rank))
    for idx in range(1, len(sequence)):
        if sequence[idx][2] < sequence[idx - 1][2]:
            problems.append(f"顺序错误：第 {sequence[idx][0]} 行「{sequence[idx][1]}」"
                            f"排在第 {sequence[idx - 1][0]} 行「{sequence[idx - 1][1]}」之后")
        if sequence[idx][2] == tail:
            seen_tail = True
        elif seen_tail:
            problems.append(f"清单外地址后面又出现了清单内地址（第 {sequence[idx][0]} 行）")

    # 3) 新行底色：A~餐种 整段 = 模板行第 A 列的颜色（含日期与类型之间的空列），
    #    总餐次/已出餐/剩余餐 三列 = 金黄；豪华行整行金黄。
    #    （2026-09-15 用户要求：不再逐列照抄模板 —— 接口读不到空格的底色。）
    base_fill = template_fills.get(plan.columns["name"], "")
    band = {c for c in (plan.columns["total"], plan.columns["served"],
                        plan.columns["left"]) if c}
    plain_range = range(1, (plan.columns["remark"] or 15) + 1)
    for change in new_changes:
        row = row_of(person_key(change.name, change.phone), change.slot)
        if not row:
            problems.append(f"新客户 {change.name}（第 {change.slot} 行）不在表里")
            continue
        luxury = str(change.meal_kind).strip() == "豪华"
        for c in plain_range:
            got = cell_fill_hex(cloud.ws.cell(row, c))
            if luxury:
                if got != f"#{GOLD}":
                    problems.append(f"{change.name} 豪华行第 {c} 列底色应为金黄，"
                                    f"实际 {got or '无'}")
                continue
            want = f"#{GOLD}" if c in band else base_fill
            if is_white(want) and is_white(got):
                continue
            if (want or "") != (got or ""):
                problems.append(f"{change.name} 第 {c} 列底色应为 {want or '无'}，"
                                f"实际 {got or '无'}")

    # 4) 新行公式：只统计真实日期列
    date_lo, date_hi = min(plan.date_cols), max(plan.date_cols)
    for change in new_changes:
        row = row_of(person_key(change.name, change.phone), change.slot)
        if not row:
            continue
        want_served = (f"=SUM({wc.column_name(date_lo)}{row}:"
                       f"{wc.column_name(date_hi)}{row})")
        got_served = str(cloud.ws.cell(row, plan.columns["served"]).value or "")
        if got_served != want_served:
            problems.append(f"{change.name} 已出餐公式应为 {want_served}，实际 {got_served!r}")
        if MARK_COL <= date_hi:
            problems.append("日期区间把协作者的标记列也算进去了")

    # 5) 协作者标记不动 + 我们的记号写对
    if cloud.ws.cell(2, MARK_COL).value != mark_before:
        problems.append(f"协作者的标记格被改动：{mark_before!r} → "
                        f"{cloud.ws.cell(2, MARK_COL).value!r}")
    got_marker = norm(cloud.ws.cell(2, MARKER_COL).value)
    if got_marker != str(plan.weekday_number):
        problems.append(f"通讯记号应为 {plan.weekday_number}，实际 {got_marker!r}")

    # 6) 行数与辅助列
    total_rows = sum(len(rows) for rows in by_person.values())
    if total_rows != len(people_before) + len(new_changes):
        problems.append(f"行数不对：原 {len(people_before)} + 新 {len(new_changes)} "
                        f"≠ 实际 {total_rows}")
    new_people = len({person_key(c.name, c.phone) for c in new_changes})
    if len(by_person) != len(people_before) + new_people:
        problems.append(f"人数不对：原 {len(people_before)} + 新客户 {new_people} "
                        f"≠ 实际 {len(by_person)}")
    if content_last_col(cloud.ws) > columns_before:
        problems.append(f"列数变多了（{columns_before} → {content_last_col(cloud.ws)}），"
                        f"辅助列可能没删干净")
    leftovers = [cloud.ws.cell(r, c).value
                 for r in range(3, cloud.ws.max_row + 1)
                 for c in range(16, cloud.ws.max_column + 1)
                 if re.fullmatch(r"\d{4}", str(cloud.ws.cell(r, c).value or ""))]
    if leftovers:
        problems.append(f"辅助列残留排序键：{leftovers[:5]}")

    # 7) 幂等：同一批第二次上传零改动（账本记住"本批已同步的本地餐次"）
    snapshot = {(r, c): cloud.ws.cell(r, c).value
                for r in range(1, cloud.ws.max_row + 1)
                for c in range(1, cloud.ws.max_column + 1)}
    plans2 = build_plan(
        cloud, local_orders={"东湖中餐": orders},
        tables={"东湖中餐": {"file_id": "FIXTURE", "drive_id": ""}},
        target=TARGET_DATE, ledger=ledger, marker_enabled=True, run_date=RUN_DATE,
        address_order={"东湖中餐": order_list}, sort_enabled=True)
    plan2 = plans2[0]
    pending2 = [c for c in plan2.changes if c.needs_write]
    if pending2:
        problems.append(f"第二次上传仍有 {len(pending2)} 条待写变更（幂等被破坏）")
    apply_plan(cloud, plans2, ledger=ledger, marker_enabled=True)
    after = {(r, c): cloud.ws.cell(r, c).value
             for r in range(1, cloud.ws.max_row + 1)
             for c in range(1, cloud.ws.max_column + 1)}
    if after != snapshot:
        changed = [k for k in after if after[k] != snapshot.get(k)]
        problems.append(f"第二次上传改动了云端内容：{changed[:5]}")

    print("\n---- 最终顺序（前 25 行）----")
    for r, addr, rank in sequence[:25]:
        print(f"  第{r:>3}行  {cloud.ws.cell(r, 1).value:<14} {addr:<8} 名次 {rank}")
    if len(sequence) > 25:
        print(f"  …共 {len(sequence)} 行")

    if problems:
        print(f"\n❌ 仿真未通过：{len(problems)} 处问题")
        for item in problems[:40]:
            print(f"  ❌ {item}")
        return 1
    twins = [c for c in new_changes if c.slot > 1]
    print(f"\n✅ 仿真通过：{len(people_before)} 人 + {len(new_changes)} 行新增"
          f"（其中同日第二单 {len(twins)} 行）；"
          f"总餐次=云端原值+本地餐次、协作者写的 {kept_zero} 个 0 原样保留、"
          f"顺序/底色/公式/记号/标记列全部符合要求，第二次上传零改动")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
